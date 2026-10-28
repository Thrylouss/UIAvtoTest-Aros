#!/usr/bin/env python3
"""
Синхронизация Excel с тест-кейсами и автотестов.

Что делает скрипт (запускается руками, когда добавили/поменяли автотесты):
  1. Читает Java-тесты и находит все методы с тегом-ID: @Tag("AUTH-001"), @Tag("CART-006") и т.д.
  2. В листе «Тест-кейсы» заполняет колонки:
       N «Статус UI-автоматизации»  → «Автоматизирован» / «Частично автоматизирован»
       O «Автотест (класс#метод)»   → storefront/cart/CartSummaryTest#cartTotalsAreConsistent
       P «Результат прогона»        → «Пройден» / «Упал» / «Пропущен» (только если указан --results)
  3. Дополнительные проверки, которых не было в Excel (ID вида AUTH-X01), выписывает в лист «Доп. проверки (авто)».
  4. Для кейсов, которые сознательно НЕ автоматизированы, записывает причину в колонку M.
  5. Дописывает в лист «Замечания» наблюдения, найденные при автоматизации (если их там ещё нет).
  6. Перед записью делает резервную копию файла в docs/test-cases/backup/.

Запуск (из корня проекта):
    python3 -m venv .venv && .venv/bin/pip install openpyxl     # один раз
    .venv/bin/python tools/sync_testcases.py --dry-run            # только показать, что изменится
    .venv/bin/python tools/sync_testcases.py                      # записать статусы автоматизации
    .venv/bin/python tools/sync_testcases.py --results target/qa-report/results.json   # + результаты прогона

Файл results.json создаётся автоматически после любого прогона тестов (mvn test …).
Примечание: после записи формулы листа «Сводка» пересчитает сам Excel при открытии файла.
"""
import argparse
import datetime
import json
import re
import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = ROOT / "docs/test-cases/Aros_Regression_Storefront_Master.xlsx"
TESTS_DIR = ROOT / "src/test/java/uz/aros/autotests/tests"

# Кейсы, где автотест покрывает только ЧАСТЬ шагов из Excel (остальное — вручную).
PARTIAL = {
    "CART-011": "Проверены перезагрузка и новая вкладка; «повторный вход» — вручную",
    "CHK-002": "Проверены способ оплаты «Наличные» и блоки; плитки баллов/кошелька — вручную",
    "CHK-012": "Проверен сценарий «наличные + самовывоз» по умолчанию",
    "PDP-001": "Проверены название, цена, «Быстрая покупка», счётчик; галерея и склады — вручную",
    "PDP-004": "Проверены кнопки «+» / «−»; ручной ввод — вручную",
    "MDL-004": "Проверен путь Samsung → Galaxy → «Аккумуляторы»",
    "CAT-004": "Проверен ввод значения «от»; поле «до» — вручную",
    "MOB-004": "Проверено отсутствие горизонтальной прокрутки на 4 страницах; наложения — визуально",
    "L10N-001": "Проверено переключение Ru → En; Oʻz — вручную",
    "L10N-003": "Проверен только раздел «Мои заказы» (статус «canceled»)",
    "HDR-003": "Проверено открытие окна и состав; фильтрация — HDR-004 вручную",
    "AUTH-014": "Проверено, что данные не показываются; но гость остаётся на пустом экране (AUTH-X03)",
}

# Почему кейс НЕ автоматизирован (пишется в колонку M «Комментарий по автоматизации»).
SKIP_REASONS = {
    "AUTH-008": "OTP по SMS отложен (решение по планированию): нужен тестовый шлюз или фиксированный код",
    "AUTH-009": "OTP по SMS отложен (решение по планированию)",
    "AUTH-010": "OTP по SMS отложен (решение по планированию)",
    "AUTH-011": "Требует SMS-код — отложено вместе с OTP",
    "HDR-004": "Каскад «область → район → магазины» нужно изучить отдельно; смена магазина влияет на остатки",
    "HDR-005": "Смена магазина меняет данные пользователя — проверять вручную",
    "HDR-009": "QR-код — визуальная проверка",
    "HOME-002": "Слайдер: нужен отдельный разбор (стрелки и точки без стабильных селекторов)",
    "CAT-002": "Структура подменю категорий требует разведки (глубина вложенности меняется)",
    "CAT-003": "Структура подменю категорий требует разведки",
    "CAT-005": "Ползунок цены — drag-and-drop, нестабилен; значения проверены в CAT-004",
    "CAT-011": "Сортировки «старые/новые» зависят от дат в данных",
    "CAT-013": "Нужен заведомо отсутствующий на складе товар — данные нестабильны",
    "CAT-014": "«Быстрый просмотр» появляется только при наведении — нужен отдельный разбор",
    "PDP-002": "Нужен товар с несколькими вариантами «Качество» — данные нестабильны",
    "PDP-003": "Остатки по складам зависят от данных",
    "PDP-005": "Нужен товар с известным малым остатком — данные нестабильны (CART-010 проверяет то же в корзине)",
    "PDP-006": "«Быстрая покупка» может создать заказ — только вручную",
    "PDP-007": "Страница «Избранное» на staging не найдена (/ru/favorites → 404) — функциональность требует уточнения",
    "PDP-008": "Галерея — визуальная проверка",
    "PDP-009": "Форма «Нашли дешевле?» отправляет данные — вручную",
    "PDP-011": "Нужен заведомо отсутствующий на складе товар",
    "PDP-013": "Нужен товар без остатка и вход пользователя",
    "FAV-001": "Страница избранного не найдена на staging (см. PDP-007)",
    "FAV-002": "Страница избранного не найдена на staging; гость видит окно входа (FAV-X01)",
    "FAV-003": "Страница избранного не найдена на staging",
    "FAV-004": "Страница избранного не найдена на staging",
    "FAV-005": "Страница избранного не найдена на staging",
    "CART-009": "Частично покрыто CART-005; массовое удаление — вручную",
    "CART-014": "Нужно менять цену/остаток через админку",
    "CHK-003": "Нужен аккаунт с балансом баллов меньше суммы заказа",
    "CHK-004": "Нужен аккаунт с достаточным балансом баллов",
    "CHK-005": "Выпадающий список доставки — требует отдельной разведки",
    "CHK-006": "Следующий шаг — выбор доставки; запускать вместе с заказом",
    "CHK-007": "Доставка BTS — требует адрес и создаёт заказ",
    "CHK-008": "Доставка в офис BTS — требует справочники и создаёт заказ",
    "CHK-009": "Доставка «До 1 км» — зависит от геолокации",
    "CHK-010": "Меняет адреса пользователя",
    "CHK-014": "СОЗДАЁТ РЕАЛЬНЫЙ ЗАКАЗ — запускать вручную (тег creates-order)",
    "CHK-015": "СОЗДАЁТ РЕАЛЬНЫЙ ЗАКАЗ — запускать вручную",
    "CHK-016": "Для проверки обязательных полей нужно нажимать «Оформить заказ»",
    "CHK-017": "Двойной клик по «Оформить заказ» — создаёт заказ(ы)",
    "CHK-018": "Нужны товары с разных складов и создаёт заказы",
    "CHK-019": "Быстрая покупка создаёт заказ",
    "CHK-020": "Нужно закрыть кошелёк через админку",
    "CHK-021": "СОЗДАЁТ ЗАКАЗ В ДОЛГ — только вручную",
    "CHK-022": "Нужен заказ в долг — только вручную",
    "CHK-023": "Нужно менять лимит долга в админке",
    "CHK-024": "Нужен просроченный долг — данные",
    "CHK-025": "Зависит от адреса и тарифа доставки",
    "CHK-026": "Карта Яндекс — нестабильна для автотестов",
    "ORD-003": "Фильтры списка — требуют разведки",
    "ORD-005": "Сверка с оформлением требует создания заказа",
    "ORD-006": "Нужна админка (смена статуса)",
    "ORD-007": "Отмена заказа меняет данные — только на специально созданном заказе",
    "ORD-008": "Нужен возврат из админки",
    "ORD-009": "Мелкая UI-проверка; можно добавить позже",
    "ORD-011": "Требует активного заказа для отмены",
    "ORD-012": "Требует активного заказа",
    "ORD-013": "Требует отмены заказа",
    "ORD-014": "Меняет способ оплаты заказа",
    "ORD-015": "Нужны завершённые и оплаченные заказы",
    "ORD-016": "Нужна админка и баллы",
    "ORD-017": "Нужен запрос продавца (админка)",
    "ORD-018": "Нужна отмена заказа и баллы",
    "PRF-002": "Редактирование данных меняет профиль пользователя — вручную",
    "PRF-003": "Требует формы редактирования (разведка)",
    "PRF-004": "Требует формы редактирования (разведка)",
    "PRF-005": "Смена телефона требует SMS",
    "PRF-006": "Меняет адреса пользователя",
    "PRF-007": "Требует формы адреса (разведка)",
    "PRF-008": "Меняет адреса пользователя",
    "PRF-009": "Меняет адреса пользователя",
    "PRF-011": "Нужен незавершённый заказ",
    "PRF-012": "Нужны долг/отрицательные баллы",
    "PRF-014": "Удаление аккаунта — необратимо",
    "PTS-001": "Страница «Баллы» (/profile/transactions) открывается (см. PRF-010); баланс и таблица — требуют разведки",
    "PTS-002": "Требует разведки",
    "PTS-003": "Нужна доставка заказа и админка",
    "PTS-004": "Нужен заказ с оплатой баллами",
    "WLT-001": "Страница кошелька открывается (см. PRF-010); сводка — требует разведки",
    "WLT-002": "Нужны заказы в долг",
    "WLT-003": "Нужны данные кошелька",
    "WLT-004": "Нужны данные кошелька",
    "WLT-005": "Нужна админка (пополнение)",
    "WLT-006": "Нужна админка (долг)",
    "NTF-001": "Страница уведомлений открывается (см. PRF-010); содержимое — требует разведки",
    "NTF-002": "Нужно создать/отменить заказ",
    "NTF-003": "Нужна админка",
    "NTF-004": "Нужна админка (спецзаказ)",
    "NTF-005": "Напоминания по таймеру — нельзя проверить автотестом",
    "NTF-006": "Требует свежих уведомлений",
    "REQ-001": "Страница запросов открывается (см. PRF-010); вкладки — требуют разведки",
    "REQ-002": "Требует разведки",
    "REQ-003": "Отправляет заявку — вручную",
    "REQ-004": "Нужна админка",
    "SPC-001": "Страница спецзаказа открывается (см. PRF-010); вкладки — требуют разведки",
    "SPC-002": "Нужна админка и создаёт заказ",
    "SPC-003": "Нужна админка",
    "SPC-004": "Таймер — нельзя проверить автотестом",
    "MST-001": "Нужен второй аккаунт (B2C) для сравнения цен",
    "MST-002": "Требует сверки цен между страницами — отдельная разведка",
    "MST-003": "Нужен второй аккаунт (B2C)",
    "MST-004": "Требует аккаунта с разными ценами",
    "MST-005": "Нужна админка (смена роли)",
    "MST-006": "Нужны данные кошелька",
    "L10N-004": "Формат цен и дат — проверка глазами на всех страницах",
    "MOB-001": "Нижнее меню мобильной версии — требует разведки",
    "MOB-002": "Фильтры на мобильном — требуют разведки",
    "MOB-003": "Оформление на мобильном создаёт заказ",
    "GEN-004": "Чат поддержки — внешний виджет",
    "GEN-005": "Время загрузки нестабильно на staging — измерять отдельно",
}

TAG_RE = re.compile(r'@Tag\("([A-Z0-9]{2,5}-X?\d{2,3})"\)')
VOID_RE = re.compile(r'\bvoid\s+(\w+)\s*\(')
DISPLAY_RE = re.compile(r'@DisplayName\("([^"]+)"\)')

RESULT_RU = {"PASSED": "Пройден", "FAILED": "Упал", "SKIPPED": "Пропущен"}

# Замечания, найденные при автоматизации (добавляются в лист «Замечания», если их ещё нет).
OBSERVATIONS = [
    ("Авторизация", "Выход из системы на странице профиля, открытой прямой ссылкой (новая вкладка), оставляет пустой экран about:blank",
     "/ru/profile → «Выйти из системы»", "Medium",
     "Похоже, выход делает «назад по истории». Автотест AUTH-X02. Через переход с главной выход работает."),
    ("Авторизация", "Гость, открывший закрытую страницу (/profile/..., /confirmation), остаётся на пустом экране — нет ни окна входа, ни редиректа",
     "/ru/profile/orders без входа", "Medium", "Данные не показываются, но UX сломан. Автотест AUTH-X03."),
    ("Авторизация", "После входа язык переключается на сохранённый в профиле (localStorage 'language'); если он отличается от языка в адресе, страница /ru показывается на английском",
     "Вход на /ru под аккаунтом с языком en", "Medium", "Адрес и интерфейс расходятся. Уточнить у разработчиков, задумано ли так."),
    ("Каталог", "У страницы каталога/поиска пустая подстановка в title вкладки: «- купить в Aros»",
     "/ru/filter?query=…", "Low", "Автотест GEN-003."),
    ("Каталог", "Товар «под заказ» показан с ценой «0 сум» при включённом «Только в наличии» (пример: Tashqi akkumulyator Xiaomi pb200szm Qora, ID 9932)",
     "/ru/filter?query=akkumulyator", "Medium", "Автотест CAT-X01."),
    ("Локализация", "Гость на /ru видит названия товаров на узбекском (запросы идут с lang=uz): «Akkumulyator Samsung…» вместо «Аккумуляторы Samsung…». У вошедшего — на русском",
     "/ru, гость", "Medium", "В запросах списка товаров параметр lang=uz независимо от языка в адресе."),
    ("Локализация", "Страница 404 на /ru показана по-английски («It seems that the product you are looking for is not available…»)",
     "/ru/favorites и другие несуществующие адреса", "Low", "Не переведена."),
    ("Локализация", "Фильтр цены в каталоге: подпись «Цена (so'm)» на узбекском в русской версии",
     "/ru/filter", "Low", "Автотест L10N-X01."),
    ("Корзина", "На странице корзины два блока «Итого:», один из них скрыт и не пересчитывается при снятии отметок (показывает старую сумму)",
     "/ru/cart", "Low", "Видимый блок считает верно. Если скрытый — мобильная панель, на телефоне сумма может быть неверной (проверить MOB-003)."),
    ("Подбор по модели", "Первая категория в списке (например, «Защитные стёкла» у Samsung Galaxy A52) при клике ничего не делает — нет ни перехода, ни сообщения",
     "Модель → Samsung → Galaxy… → первая категория", "Low", "Ожидается переход со списком или сообщение «нет товаров»."),
    ("API", "Запросы с некорректными значениями: sidebar-filters отвечает 500 (?brands=undefined&product_models=undefined); …/warehouses/null/workers/ — 404; …/product-variants/quantities/ — 400",
     "/ru/filter, главная (вошедший пользователь)", "Medium", "Фронт отправляет undefined/null. Нужны проверки на стороне фронта и бэка. Автотест GEN-002 ловит 5xx."),
    ("Избранное", "Страница избранного не найдена: /ru/favorites возвращает 404; у вошедшего клик по «сердцу» в шапке ничего не открывает",
     "Шапка → «сердце»", "Medium", "Уточнить у разработчиков: функциональность не реализована или маршрут другой. Автотесты FAV-* не написаны."),
]


def scan_tests():
    """Возвращает {caseId: [(путь_класса, метод, DisplayName, [теги])]}."""
    found = {}
    for path in sorted(TESTS_DIR.rglob("*.java")):
        text = path.read_text(encoding="utf-8")
        rel = path.relative_to(TESTS_DIR).with_suffix("")
        for m in VOID_RE.finditer(text):
            start = text.rfind("\n\n", 0, m.start())
            block = text[start if start >= 0 else 0:m.start()]
            ids = TAG_RE.findall(block)
            if not ids:
                continue
            display = DISPLAY_RE.search(block)
            tags = re.findall(r'@Tag\("([^"]+)"\)', block)
            for cid in dict.fromkeys(ids):
                found.setdefault(cid, []).append((str(rel).replace("\\", "/"), m.group(1),
                                                  display.group(1) if display else "", tags))
    return found


def load_results(path):
    """Сводит результаты по ID: если хоть один запуск кейса упал — «Упал»."""
    if not path:
        return {}
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    order = {"FAILED": 3, "PASSED": 2, "SKIPPED": 1}
    out = {}
    for r in data:
        cid = r.get("caseId")
        if cid and (cid not in out or order[r["status"]] > order[out[cid]["status"]]):
            out[cid] = r
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", default=str(DEFAULT_XLSX))
    ap.add_argument("--results", help="путь к results.json (target/qa-report/results.json)")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    xlsx = Path(args.xlsx)
    tests = scan_tests()
    results = load_results(args.results)
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["Тест-кейсы"]

    header = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    col_status = header["Статус UI-автоматизации"]
    col_test = header["Автотест (класс#метод)"]
    col_result = header["Результат прогона"]
    col_comment_auto = next(c for h, c in header.items() if h.startswith("Комментарий по автоматиз"))

    changed = 0
    covered = set()
    for row in range(2, ws.max_row + 1):
        cid = ws.cell(row, 1).value
        if not cid:
            continue
        cid = str(cid).strip()
        new = {}
        if cid in tests:
            covered.add(cid)
            new[col_status] = "Частично автоматизирован" if cid in PARTIAL else "Автоматизирован"
            new[col_test] = "\n".join(sorted({f"{c}#{m}" for c, m, _, _ in tests[cid]}))
            if cid in PARTIAL:
                new[col_comment_auto] = PARTIAL[cid]
            if cid in results:
                new[col_result] = RESULT_RU[results[cid]["status"]]
        elif cid in SKIP_REASONS:
            new[col_comment_auto] = SKIP_REASONS[cid]
        for col, val in new.items():
            if ws.cell(row, col).value != val:
                print(f"  {cid:9} {str(ws.cell(1, col).value)[:26]:26} ← {str(val)[:70]!r}")
                ws.cell(row, col).value = val
                changed += 1

    extras = {cid: refs for cid, refs in tests.items() if re.match(r"^[A-Z0-9]+-X\d+$", cid)}
    sheet_name = "Доп. проверки (авто)"
    if extras:
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ex = wb.create_sheet(sheet_name)
        ex.append(["ID", "Название", "Автотест (класс#метод)", "Известный дефект?", "Результат прогона"])
        for cid, refs in sorted(extras.items()):
            c, m, display, tags = refs[0]
            title = re.sub(r"^[A-Z0-9]+-X?\d+\s*[·:\-–]\s*", "", display)
            res = RESULT_RU[results[cid]["status"]] if cid in results else "Не выполнен"
            ex.append([cid, title, f"{c}#{m}", "Да" if "known-defect" in tags else "Нет", res])
        for col, width in zip("ABCDE", (11, 80, 55, 18, 18)):
            ex.column_dimensions[col].width = width
        for cell in ex[1]:
            cell.font = Font(bold=True)
        print(f"  лист «{sheet_name}»: {len(extras)} дополнительных проверок")

    obs = wb["Замечания"]
    existing = {str(obs.cell(r, 3).value) for r in range(2, obs.max_row + 1)}
    footer_row = next((r for r in range(2, obs.max_row + 1)
                       if obs.cell(r, 2).value and str(obs.cell(r, 2).value).startswith("Замечания найдены")), None)
    footer_text = obs.cell(footer_row, 2).value if footer_row else None
    last_no = max((int(obs.cell(r, 1).value) for r in range(2, obs.max_row + 1)
                   if str(obs.cell(r, 1).value).isdigit()), default=0)
    first_free = (footer_row - 1) if footer_row else obs.max_row + 1
    new_rows = [o for o in OBSERVATIONS if o[1] not in existing]
    if new_rows:
        if footer_row:
            obs.cell(footer_row, 2).value = None
        row = first_free
        for cat, desc, where, sev, comment in new_rows:
            last_no += 1
            for col, val in enumerate([last_no, cat, desc, where, sev, comment, "Новое"], start=1):
                obs.cell(row, col).value = val
                obs.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
        if footer_text:
            obs.cell(row + 1, 2).value = footer_text
        print(f"  лист «Замечания»: добавлено {len(new_rows)}")

    print(f"Кейсов Excel с автотестами: {len(covered)}; изменено ячеек: {changed}")
    unknown = sorted(c for c in tests if c not in covered and c not in extras)
    if unknown:
        print("ВНИМАНИЕ: в тестах есть ID, которых нет в Excel:", ", ".join(unknown))

    if args.dry_run:
        print("--dry-run: файл не изменён")
        return
    backup_dir = xlsx.parent / "backup"
    backup_dir.mkdir(exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    shutil.copy2(xlsx, backup_dir / f"{xlsx.stem}-{stamp}{xlsx.suffix}")
    wb.save(xlsx)
    print(f"Сохранено: {xlsx}  (резервная копия: backup/{xlsx.stem}-{stamp}{xlsx.suffix})")


if __name__ == "__main__":
    sys.exit(main())
